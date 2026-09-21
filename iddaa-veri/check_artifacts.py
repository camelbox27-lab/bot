import openpyxl
from datetime import datetime

def parse_date(d):
    if not d:
        return None
    s = str(d).strip()
    for fmt in ['%d.%m.%Y', '%m/%d/%y', '%Y-%m-%d', '%d/%m/%Y']:
        try:
            return datetime.strptime(s, fmt)
        except:
            pass
    return None

files = [
    ('output/artifacts/eksik3/eksik-2017-08-2018-03/eksik_2017_08_2018_03.xlsx', 'Eksik 2017-08~2018-03'),
    ('output/artifacts/eksik2/eksik-2018-04-2018-11/eksik_2018_04_2018_11.xlsx', 'Eksik 2018-04~2018-11'),
    ('output/artifacts/eksik1/eksik-2018-12-2019-07/eksik_2018_12_2019_07.xlsx', 'Eksik 2018-12~2019-07'),
    ('output/artifacts/guncel/guncel-2026-03-2026-05/guncel_2026_03_2026_05.xlsx', 'Guncel 2026-03~2026-05'),
]

print('Workflow Sonuclari:')
print('='*60)
for fpath, name in files:
    wb = openpyxl.load_workbook(fpath, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    data = rows[1:]
    dates = sorted([parse_date(r[2]) for r in data if r[2] and parse_date(r[2])])
    fmt = '%Y-%m-%d'
    if dates:
        print(name)
        print('  Mac sayisi : ' + str(len(data)))
        print('  Ilk tarih  : ' + dates[0].strftime(fmt))
        print('  Son tarih  : ' + dates[-1].strftime(fmt))
    wb.close()
    print()

# Patch dosyasini da kontrol et
print('Patch Sonucu (iddaagecmismaclar_patched.xlsx):')
print('='*60)
wb = openpyxl.load_workbook('output/artifacts/patch/patch-clean-excel/iddaagecmismaclar_patched.xlsx', read_only=True)
ws = wb.active
rows_p = list(ws.iter_rows(values_only=True))
data_p = rows_p[1:]
print('Header:', rows_p[0][:5], '...')
print('Toplam satir:', len(data_p))
dates_p = sorted([parse_date(r[2]) for r in data_p if len(r)>2 and r[2] and parse_date(r[2])])
if dates_p:
    print('Ilk tarih  : ' + dates_p[0].strftime(fmt))
    print('Son tarih  : ' + dates_p[-1].strftime(fmt))
wb.close()
