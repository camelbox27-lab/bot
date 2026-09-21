import openpyxl
from collections import Counter

fpath = 'output/artifacts/patch/patch-clean-excel/iddaagecmismaclar_patched.xlsx'
wb = openpyxl.load_workbook(fpath, read_only=True)
ws = wb.active

header = [cell.value for cell in ws[1]]
iy_col = header.index('IY Skor')
ms_col = header.index('MS Skor')

vals = Counter()
examples = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    v = row[iy_col]
    s = str(v) if v is not None else 'NONE'
    vals[s] += 1
    if s not in examples:
        examples[s] = (row[0], row[1], row[2], row[ms_col])  # ev, dep, tarih, ms

wb.close()

print("IY Skor deger dagilimi (top 30):")
for v, c in vals.most_common(30):
    ex = examples[v]
    print(f"  {c:6d}x  {repr(v):30s}  ornek: {ex[0]} vs {ex[1]} ({ex[2]}) MS={ex[3]}")
