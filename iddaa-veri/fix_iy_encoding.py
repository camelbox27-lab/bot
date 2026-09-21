import openpyxl
import re
from openpyxl.styles import PatternFill, Font, Alignment

SRC = 'output/artifacts/patch/patch-clean-excel/iddaagecmismaclar_patched.xlsx'
DST = 'output/iddaagecmismaclar_final.xlsx'

print("Excel okunuyor (buyuk dosya)...")
wb = openpyxl.load_workbook(SRC)
ws = wb.active

header = [cell.value for cell in ws[1]]
iy_col = header.index('IY Skor') + 1  # 1-based

print(f"Toplam satir: {ws.max_row - 1:,}")
print("IY encoding fix uygulanıyor...")

score_re = re.compile(r'(\d+)\s*[-:]\s*(\d+)')
fixed_extracted = 0
fixed_cleared = 0

for row in ws.iter_rows(min_row=2):
    cell = row[iy_col - 1]
    v = cell.value
    if not v:
        continue
    s = str(v).strip()

    # Zaten dogru format (rakam ile basliyor): dokunma
    if s and s[0].isdigit():
        continue
    # "-" veya bos: dokunma
    if s in ('-', ''):
        continue

    # "?Y 1-0" gibi - skoru cikart
    m = score_re.search(s)
    if m:
        cell.value = f"{m.group(1)}-{m.group(2)}"
        fixed_extracted += 1
    else:
        # "?Y" tek basina - skor yok, temizle
        cell.value = '-'
        fixed_cleared += 1

print(f"  Skor cikartildi (encoding fix): {fixed_extracted:,}")
print(f"  Temizlendi (skor yok):          {fixed_cleared:,}")
print(f"\nKaydediliyor -> {DST}")
wb.save(DST)
print("Kaydedildi.")

# Dogrulama
print("\nDogrulama yapiliyor...")
wb2 = openpyxl.load_workbook(DST, read_only=True)
ws2 = wb2.active
header2 = [c.value for c in ws2[1]]
iy2 = header2.index('IY Skor')
ms2 = header2.index('MS Skor')

still_bozuk = 0
ms_eksik = 0
iy_temiz = 0
for row in ws2.iter_rows(min_row=2, values_only=True):
    iy_v = str(row[iy2]).strip() if row[iy2] else ''
    ms_v = str(row[ms2]).strip() if row[ms2] else ''
    if iy_v and iy_v not in ('-', '') and not iy_v[0].isdigit():
        still_bozuk += 1
    if ms_v in ('-', '', 'None'):
        ms_eksik += 1
    if iy_v and iy_v[0].isdigit():
        iy_temiz += 1

wb2.close()
print(f"  IY Skor temiz (X-X format): {iy_temiz:,}")
print(f"  IY Skor hala bozuk        : {still_bozuk}")
print(f"  MS Skor eksik             : {ms_eksik:,}")
print("\nTAMAM.")
